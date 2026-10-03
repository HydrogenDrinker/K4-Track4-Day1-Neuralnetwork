"""results_table.py — Serialization of results to JSON and populating experiments.xlsx workbook.

Maintains all Excel formulas intact across Experiments, Seeds, and Summary sheets.
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra JSON."""
    exp_id = result["cfg"]["exp_id"]
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{exp_id}.json"

    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return str(out_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict."""
    p = Path(results_dir)
    if not p.exists():
        return []

    results = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            results.append(json.load(fp))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng Experiments."""
    cfg = result["cfg"]
    summary = result.get("summary", {})
    exp_id = cfg["exp_id"]

    eval_acc = None
    eval_macro_f1 = None
    if eval_scores:
        eval_acc = eval_scores.get("accuracy")
        eval_macro_f1 = eval_scores.get("macro_f1")

    hidden_str = ", ".join(str(h) for h in cfg.get("hidden", []))

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", ""),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_acc,
        "eval_macro_f1": eval_macro_f1,
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               summary_notes: dict | None = None,
               baseline_seeds: list[str] | None = None) -> None:
    """Điền các dòng vào template và lưu thành file out_path."""
    wb = openpyxl.load_workbook(template_path)
    ws_exp = wb["Experiments"]

    # Đọc headers từ dòng 1
    headers = [ws_exp.cell(1, col).value for col in range(1, ws_exp.max_column + 1)]

    # Ghi dữ liệu vào sheet Experiments (chỉ ghi cột 1..29, giữ nguyên công thức cột 30..33)
    for r_idx, row in enumerate(rows, start=2):
        for col_idx in range(1, min(30, len(headers) + 1)):
            header_name = headers[col_idx - 1]
            if header_name in row:
                val = row[header_name]
                ws_exp.cell(r_idx, col_idx, value=val)

    # Cập nhật sheet Seeds
    if "Seeds" in wb.sheetnames and baseline_seeds:
        ws_seeds = wb["Seeds"]
        for i, s_id in enumerate(baseline_seeds):
            ws_seeds.cell(row=2 + i, column=1, value=s_id)

    # Cập nhật nhận xét vào sheet Summary
    if "Summary" in wb.sheetnames and summary_notes:
        ws_summary = wb["Summary"]
        for r in range(2, ws_summary.max_row + 1):
            grp = ws_summary.cell(r, 1).value
            if grp in summary_notes:
                ws_summary.cell(r, 8, value=summary_notes[grp])

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Đã lưu bảng thí nghiệm với {len(rows)} dòng vào {out_path}")
